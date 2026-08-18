"""
Builds docs/Tracksure-Driver-Demo.pptx — the demo deck for the driver PWA.

Six 16:9 slides: the flow at a glance, then the three capture moments the whole
data pipeline depends on (waybill, fuel receipt, declaration), and what each one
replaces on paper.

BRAND.md composition: white slides with ink text and exactly one orange accent;
the title slide uses the charcoal --grad-side ground, which is the one place the
original horizontal lockup (white wordmark) belongs.
"""
import sys
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

S = sys.argv[1]
OUT = sys.argv[2]

BRAND = RGBColor(0xE0, 0x75, 0x1A)
BRAND_LIGHT = RGBColor(0xFF, 0xA3, 0x00)
INK = RGBColor(0x30, 0x30, 0x30)
INK_SOFT = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DIM = RGBColor(0xC9, 0xC7, 0xC4)
FONT = "Arial"   # see the note in build_quickstart.py — brand fonts are not installed

W, H = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]


def textbox(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    return tf


def write(tf, text, *, size, color, bold=False, space_after=0, first=False, line=1.15):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.line_spacing = line
    p.space_after = Pt(space_after)
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = FONT
    return p


def bullet(tf, text, *, first=False, color=INK_SOFT, size=17):
    p = write(tf, f"—   {text}", size=size, color=color, space_after=11, first=first, line=1.25)
    return p


def phone(slide, img, *, right_edge=Inches(12.5), top=Inches(0.95), height=Inches(5.9)):
    """Drop a screenshot in at a fixed height, right-aligned."""
    pic = slide.shapes.add_picture(f"{S}/quick-start-assets/{img}.png", 0, top, height=height)
    pic.left = Emu(int(right_edge) - pic.width)
    return pic


def accent(slide, x, y, w=Inches(1.5)):
    slide.shapes.add_picture(f"{S}/quick-start-assets/slide-bar.png", x, y, width=w, height=Inches(0.09))


def content_slide(title, kicker, bullets, img, *, imgs=None):
    s = prs.slides.add_slide(BLANK)
    accent(s, Inches(0.9), Inches(0.95))

    tf = textbox(s, Inches(0.9), Inches(1.25), Inches(6.2), Inches(1.2))
    write(tf, kicker, size=15, color=BRAND, bold=True, space_after=8, first=True)
    write(tf, title, size=38, color=INK, bold=True, line=1.05)

    tf2 = textbox(s, Inches(0.9), Inches(3.15), Inches(6.1), Inches(3.4))
    for i, b in enumerate(bullets):
        bullet(tf2, b, first=(i == 0))

    if imgs:
        # Two phones side by side (the pair of declarations).
        # Sized to fit: two phones plus the gap must land inside 7.1in..12.9in.
        x = Inches(7.15)
        for name in imgs:
            pic = s.shapes.add_picture(f"{S}/quick-start-assets/{name}.png", x, Inches(1.35), height=Inches(4.5))
            x = Emu(int(x) + pic.width + int(Inches(0.3)))
    else:
        phone(s, img)
    return s


# ------------------------------------------------------------- 1. title ----
s = prs.slides.add_slide(BLANK)
s.shapes.add_picture(f"{S}/quick-start-assets/slide-bg-dark.png", 0, 0, width=W, height=H)
s.shapes.add_picture(f"{S}/quick-start-assets/assets-logo-ondark.png", Inches(1.0), Inches(1.15), width=Inches(3.5))

tf = textbox(s, Inches(1.0), Inches(2.5), Inches(6.6), Inches(3.0))
write(tf, "Tracksure Driver", size=54, color=WHITE, bold=True, space_after=14, first=True, line=1.0)
write(tf, "The paper trail, captured at the roadside", size=26, color=BRAND_LIGHT, bold=True, space_after=26)
write(tf, "Every waybill, fuel receipt and signature enters the system here —\n"
          "from a phone, at the moment it happens.", size=19, color=DIM, line=1.35)

# Phones peeking in from the right.
for i, name in enumerate(("02-home-blocking", "07-waybill-post")):
    s.shapes.add_picture(f"{S}/quick-start-assets/{name}.png", Inches(7.85 + i * 2.4), Inches(1.6), height=Inches(4.4))

# -------------------------------------------------------------- 2. flow ----
s = prs.slides.add_slide(BLANK)
accent(s, Inches(0.9), Inches(0.85))
tf = textbox(s, Inches(0.9), Inches(1.15), Inches(11.5), Inches(1.2))
write(tf, "THE WHOLE JOB", size=15, color=BRAND, bold=True, space_after=8, first=True)
write(tf, "Five taps, start to finish", size=38, color=INK, bold=True, line=1.05)

steps = [
    ("01-login", "Sign in"),
    ("02-home-blocking", "See the trip"),
    ("04-declaration-pretrip", "Sign to load"),
    ("07-waybill-post", "Photograph"),
    ("10-fuel-sent", "Log fuel"),
]
# 5 phones + 4 gaps must fit the 11.53in between the margins.
x = Inches(0.97)
for name, label in steps:
    pic = s.shapes.add_picture(f"{S}/quick-start-assets/{name}.png", x, Inches(2.5), height=Inches(3.25))
    cap = textbox(s, Emu(int(x) - int(Inches(0.2))), Inches(6.0), Emu(pic.width + int(Inches(0.4))), Inches(0.5))
    p = write(cap, label, size=16, color=INK, bold=True, first=True)
    p.alignment = PP_ALIGN.CENTER
    x = Emu(int(x) + pic.width + int(Inches(0.3)))

# ----------------------------------------------------------- 3. waybill ----
content_slide(
    "The signed photo IS the delivery",
    "CAPTURE MOMENT 1",
    [
        "Camera-first. The driver photographs the waybill the customer signed.",
        "GPS and the time are stamped on it automatically, and checked against the destination.",
        "It reaches Operations while the truck is still at the gate — not days later in a folder.",
        "No OTP, no signature pad, nothing else to carry.",
    ],
    "07-waybill-post",
)

# -------------------------------------------------------------- 4. fuel ----
content_slide(
    "A fuel receipt that files itself",
    "CAPTURE MOMENT 2",
    [
        "The driver photographs the receipt and picks the station. That is the whole job.",
        "Credit or cash resolves from the station record — the driver is never asked, and never guesses.",
        "Litres land against the trip's fuel estimate, so variance is visible the same day.",
        "The photo is compressed on the phone first: drivers pay for their own data.",
    ],
    "10-fuel-sent",
)

# ------------------------------------------------------ 5. declarations ----
content_slide(
    "A PIN is a signature",
    "CAPTURE MOMENT 3",
    [
        "The driver signs the journey plan with their own PIN, before departure and on return.",
        "Both are timestamped and audit-logged. This is the digital signature on an HSE document.",
        "The pre-departure one gates loading — and the app says so, in those words.",
        "The return one never blocks the next load. A reminder, not a barrier.",
    ],
    None,
    imgs=("04-declaration-pretrip", "13-declaration-return"),
)

# ------------------------------------------------------- 6. kills paper ----
s = prs.slides.add_slide(BLANK)
accent(s, Inches(0.9), Inches(0.85))
tf = textbox(s, Inches(0.9), Inches(1.15), Inches(11.5), Inches(1.2))
write(tf, "WHY IT MATTERS", size=15, color=BRAND, bold=True, space_after=8, first=True)
write(tf, "What stops being paper", size=38, color=INK, bold=True, line=1.05)

left = textbox(s, Inches(0.9), Inches(2.7), Inches(5.5), Inches(4.0))
write(left, "TODAY", size=15, color=INK_SOFT, bold=True, space_after=14, first=True)
for t in (
    "Waybills photocopied back at the yard, days after delivery.",
    "Fuel receipts in a pocket. Some arrive, some do not.",
    "Journey plan signed on a form, filed in a folder.",
    "Ops finds out what happened when the truck comes back.",
):
    bullet(left, t)

right = textbox(s, Inches(7.0), Inches(2.7), Inches(5.5), Inches(4.0))
write(right, "WITH TRACKSURE", size=15, color=BRAND, bold=True, space_after=14, first=True)
for t in (
    "Photo, GPS and timestamp, the moment the customer signs.",
    "Receipt tagged to the station and the trip, litres counted.",
    "PIN signature, timestamped and audit-logged against the plan.",
    "Ops sees it while the truck is still on the road.",
):
    bullet(right, t, color=INK)

prs.save(OUT)
print(f"wrote {OUT} — {len(prs.slides.__iter__.__self__._sldIdLst)} slides")
